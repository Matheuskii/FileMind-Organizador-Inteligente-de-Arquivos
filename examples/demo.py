"""Execute da raiz: python -m examples.demo [--live].

Padrão: classificador de exemplo com respostas fixas, sem API ou chave.
--live: envia conteúdo fictício à Groq para classificação real.
"""
from __future__ import annotations

import argparse
import logging
import os
import tempfile
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from dotenv import load_dotenv
from PIL import Image, ImageDraw
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from main import Classifier, Organizer, digest, validate_classification


class ExampleClassifier:
    """Respostas fixas por fixture; não representa uma IA nem analisa arquivos reais."""

    def classify(self, evidence):
        import json

        category, subcategory, title = {
            "download_001.TXT": ("Financeiro", "NotasFiscais", "NotaFiscal_LojaExemplo"),
            "arquivo_002.PDF": ("Estudos", "Apostilas", "Apostila_Python"),
            "documento_003.docx": ("Trabalho", "Contratos", "Contrato_EmpresaExemplo"),
            "imagem_004.PNG": ("Midia", "CapturasDeTela", "Captura_FileMind"),
        }[evidence["original_name"]]
        # Exercita a mesma validação de JSON usada pela integração real.
        return validate_classification(json.dumps({
            "category": category, "subcategory": subcategory, "title": title,
            "document_date": "2026-10-16", "confidence": .99,
        }))


def create_samples(folder: Path) -> list[Path]:
    """Gera quatro arquivos pequenos, inclusive PDF com texto extraível."""
    folder.mkdir(parents=True, exist_ok=True)
    txt = folder / "download_001.TXT"
    txt.write_text("NOTA FISCAL - Loja Exemplo\nData: 16/10/2026\nTeclado USB: R$ 120,00\n",
                   encoding="utf-8")

    pdf = folder / "arquivo_002.PDF"
    writer = PdfWriter()
    page = writer.add_blank_page(width=595, height=842)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 16 Tf 50 780 Td (Apostila de Python - 16/10/2026) Tj "
                    b"0 -28 Td (Variaveis, funcoes e automacao de arquivos.) Tj ET")
    page[NameObject("/Contents")] = stream
    writer.add_metadata({"/Title": "Apostila de Python"})
    writer.write(pdf)

    docx = folder / "documento_003.docx"
    document = Document()
    document.add_heading("Contrato de trabalho - Empresa Exemplo", 0)
    document.add_paragraph("Data: 16/10/2026. Documento fictício para teste do FileMind.")
    document.save(docx)

    png = folder / "imagem_004.PNG"
    image = Image.new("RGB", (640, 240), "white")
    ImageDraw.Draw(image).text((30, 80), "FileMind - captura ficticia - 16/10/2026", fill="black")
    image.save(png)
    return [txt, pdf, docx, png]


def run_demo(base: Path, live: bool = False) -> tuple[Path, list[Path]]:
    # Sempre cria subpasta única, sem limpar nem sobrescrever execuções anteriores.
    base = base.expanduser().resolve()
    base.mkdir(parents=True, exist_ok=True)
    run = Path(tempfile.mkdtemp(prefix="run-", dir=base))
    sources = create_samples(run / "entrada")
    args = SimpleNamespace(
        watch=run / "entrada", output=run / "organizados", state=run / "state.sqlite3",
        max_mb=30, max_chars=12000, pdf_pages=8, ocr=False, ocr_lang="por+eng",
        threshold=.75, ready_timeout=30, apply=True, mode="copy",
    )
    classifier = (Classifier(os.getenv("GROQ_MODEL", "llama-3.3-70b-versatile"), False)
                  if live else ExampleClassifier())
    organizer = Organizer(args, classifier)
    for source in sources:
        organizer.process(source)
    targets = sorted(args.output.rglob("*"))
    targets = [path for path in targets if path.is_file()]
    # A saída da IA pode mudar; os invariantes de integridade são os mesmos.
    for source in sources:
        if not any(target.suffix == source.suffix and digest(target) == digest(source) for target in targets):
            raise AssertionError(f"Extensão ou conteúdo divergente: {source.name}")
    if len(targets) != len(sources):
        raise AssertionError("Quantidade incorreta de arquivos organizados")
    return run, targets


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--live", action="store_true", help="usa API real da Groq")
    parser.add_argument("--workspace", type=Path, default=Path("demo-data"))
    args = parser.parse_args()
    load_dotenv()
    if args.live and os.getenv("GROQ_API_KEY", "") in ("", "coloque_sua_chave_aqui"):
        parser.error("preencha GROQ_API_KEY no .env para usar --live")
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    print("Modo: Groq REAL" if args.live else "Modo: respostas FIXAS de exemplo, sem API", flush=True)
    run, targets = run_demo(args.workspace, args.live)
    for target in targets:
        print(target.relative_to(run))
    print(f"OK: 4 arquivos copiados; extensões e SHA-256 preservados. Pasta: {run}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
