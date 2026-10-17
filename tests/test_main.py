import json
import threading
import time
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from docx import Document
from PIL import Image
from pypdf import PdfWriter

from main import (Classification, Classifier, Organizer, destination_for, digest, extract,
                  fingerprint, publish_copy, safe_component, validate_classification, wait_ready)


def result(**changes):
    payload = dict(category="Financeiro", subcategory="NotasFiscais", title="NotaFiscal_Loja",
                   document_date="2026-10-16", confidence=.95)
    payload.update(changes)
    return Classification(**payload)


@pytest.mark.parametrize("extension", [".pdf", ".PDF", ".JpEg", ".docx", ".TXT"])
def test_original_extension_preserved(tmp_path, extension):
    target = destination_for(tmp_path / ("download" + extension), tmp_path, result(), .75)
    assert target.name == "2026-10-16_NotaFiscal_Loja" + extension


@pytest.mark.parametrize("changes", [dict(category="../../escape"), dict(subcategory="Contratos"),
                                        dict(document_date="2026-02-30"), dict(confidence=2),
                                        dict(confidence="0.9"), dict(extra="forbidden"), dict(title="")])
def test_reject_invalid_model_output(changes):
    payload = result().model_dump()
    payload.update(changes)
    with pytest.raises(ValueError):
        validate_classification(json.dumps(payload))


def test_sanitize_and_review(tmp_path):
    assert safe_component("../../CON") == "Arquivo_CON"
    path = destination_for(tmp_path / "a.PDF", tmp_path,
                           result(title="../../Nota: Loja.pdf", document_date=None, confidence=.2), .75)
    assert path.relative_to(tmp_path).parts == (
        "Outros", "Revisar", "SemData_Nota_Loja_pdf.PDF")


@pytest.mark.parametrize("move", [True, False])
def test_collision_preserves_bytes(tmp_path, move):
    source = tmp_path / "original.TXT"
    source.write_bytes(b"documento original")
    target = tmp_path / "organized/a.TXT"
    target.parent.mkdir()
    target.write_bytes(b"nao sobrescrever")
    final = publish_copy(source, target, fingerprint(source), digest(source), move)
    assert target.read_bytes() == b"nao sobrescrever"
    assert final.name == "a_001.TXT"
    assert final.read_bytes() == b"documento original"
    assert source.exists() is not move


def test_changed_source_never_deleted(tmp_path):
    source = tmp_path / "a.txt"
    source.write_text("primeiro")
    before, sha = fingerprint(source), digest(source)
    source.write_text("outro conteúdo")
    target = tmp_path / "out/a.txt"
    with pytest.raises(ValueError):
        publish_copy(source, target, before, sha, True)
    assert source.exists()
    assert not target.exists()


def test_extract_formats(tmp_path):
    txt = tmp_path / "a.txt"
    txt.write_text("Olá mundo", encoding="utf-8")
    assert extract(txt, 100, 2, False, "por")["text"] == "Olá mundo"
    docx = tmp_path / "a.docx"
    document = Document()
    document.add_paragraph("Contrato de trabalho")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Empresa Exemplo"
    document.save(docx)
    assert "Empresa Exemplo" in extract(docx, 1000, 2, False, "por")["text"]
    pdf = tmp_path / "a.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.add_metadata({"/Title": "Apostila"})
    writer.write(pdf)
    assert extract(pdf, 1000, 2, False, "por")["metadata"]["/Title"] == "Apostila"
    png = tmp_path / "a.png"
    Image.new("RGB", (10, 20)).save(png)
    evidence = extract(png, 1000, 2, False, "por")
    assert evidence["metadata"]["height"] == 20
    assert not evidence["has_text"]


def test_ready_timeout_and_zero_byte(tmp_path):
    source = tmp_path / "a.txt"
    source.touch()
    with pytest.raises(TimeoutError):
        wait_ready(source, threading.Event(), timeout=.02, interval=.001)
    source.write_text("pronto")
    assert wait_ready(source, threading.Event(), timeout=.1, interval=.001) == fingerprint(source)


def test_symlink_source_rejected(tmp_path):
    source = tmp_path / "real.txt"
    source.write_text("real")
    link = tmp_path / "link.txt"
    try:
        link.symlink_to(source)
    except OSError:
        pytest.skip("symlink indisponível sem privilégio no Windows")
    with pytest.raises(ValueError):
        fingerprint(link)


def test_model_validation_retry_and_schema_modes(monkeypatch):
    monkeypatch.setenv("GROQ_API_KEY", "fake-test-key")
    monkeypatch.setattr("main.Groq", Mock(return_value=Mock()))
    classifier = Classifier("llama-3.3-70b-versatile", False)
    create = Mock(side_effect=[
        SimpleNamespace(choices=[SimpleNamespace(finish_reason="stop", message=SimpleNamespace(content='{}'))]),
        SimpleNamespace(choices=[SimpleNamespace(finish_reason="stop", message=SimpleNamespace(
            content=result().model_dump_json()))])])
    classifier.client.chat.completions.create = create
    assert classifier.classify({"text": "exemplo"}).category == "Financeiro"
    assert create.call_count == 2
    assert create.call_args.kwargs["response_format"] == {"type": "json_object"}
    classifier.strict_server = True
    create.side_effect = None
    create.return_value = SimpleNamespace(choices=[SimpleNamespace(
        finish_reason="stop", message=SimpleNamespace(content=result().model_dump_json()))])
    classifier.classify({})
    assert create.call_args.kwargs["response_format"]["json_schema"]["strict"] is True


def test_pipeline_simulation_apply_and_idempotence(tmp_path, monkeypatch):
    watch, output = tmp_path / "in", tmp_path / "out"
    watch.mkdir()
    source = watch / "download.TXT"
    source.write_text("Nota fiscal Loja Exemplo 16/10/2026")
    args = SimpleNamespace(state=tmp_path / "state.sqlite3", watch=watch, output=output,
                           ready_timeout=1, max_mb=30, max_chars=12000, pdf_pages=8, ocr=False,
                           ocr_lang="por", threshold=.75, apply=False, mode="copy")
    classifier = Mock()
    classifier.classify.return_value = result()
    organizer = Organizer(args, classifier)
    monkeypatch.setattr("main.wait_ready", lambda path, *_: fingerprint(path))
    organizer.submit(source)
    organizer.submit(source)
    assert organizer.pending.qsize() == 1
    organizer.process(source)
    assert not output.exists()
    args.apply = True
    organizer.process(source)
    organizer.process(source)
    assert len(list(output.rglob("*.TXT"))) == 1
    assert classifier.classify.call_count == 2
    assert source.exists()


def test_failed_ai_keeps_original(tmp_path, monkeypatch):
    source = tmp_path / "source.txt"
    source.write_text("Documento")
    args = SimpleNamespace(state=tmp_path / "state.sqlite3", watch=tmp_path,
                           output=tmp_path / "out", ready_timeout=1, max_mb=30,
                           max_chars=12000, pdf_pages=8, ocr=False, ocr_lang="por",
                           threshold=.75, apply=True, mode="move")
    classifier = Mock()
    classifier.classify.side_effect = ValueError("resposta inválida")
    monkeypatch.setattr("main.wait_ready", lambda path, *_: fingerprint(path))
    with pytest.raises(ValueError):
        Organizer(args, classifier).process(source)
    assert source.read_text() == "Documento"
    assert not args.output.exists()


def test_real_watchdog_event_processes_new_file(tmp_path, monkeypatch):
    from watchdog.observers import Observer

    watch, output = tmp_path / "in", tmp_path / "out"
    watch.mkdir()
    args = SimpleNamespace(state=tmp_path / "state.sqlite3", watch=watch, output=output,
                           ready_timeout=1, max_mb=30, max_chars=12000, pdf_pages=8, ocr=False,
                           ocr_lang="por", threshold=.75, apply=True, mode="copy")
    classifier = Mock()
    classifier.classify.return_value = result()
    organizer = Organizer(args, classifier)
    monkeypatch.setattr("main.wait_ready", lambda path, *_: fingerprint(path))
    worker = threading.Thread(target=organizer.worker)
    observer = Observer()
    observer.schedule(organizer, str(watch), recursive=False)
    worker.start()
    observer.start()
    try:
        source = watch / "incoming.TXT"
        source.write_text("Nota fiscal Loja Exemplo")
        deadline = time.monotonic() + 5
        while not list(output.rglob("*.TXT")) and time.monotonic() < deadline:
            time.sleep(.05)
        organizer.pending.join()
        files = list(output.rglob("*.TXT"))
        assert len(files) == 1
        assert files[0].read_bytes() == source.read_bytes()
        assert organizer.failures == 0
    finally:
        observer.stop()
        observer.join()
        organizer.stop.set()
        worker.join(timeout=3)
